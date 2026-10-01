
"""
FILE 5: evaluate.py
PURPOSE: Evaluate models — saves ALL metrics to Excel file
         Every time you run → new row added to Excel
         Excel has 5 sheets → ready for report tables!
RUN: python evaluate.py
"""

import os
import numpy as np
import pandas as pd
import pickle
import matplotlib.pyplot as plt
import seaborn as sns
from datetime import datetime
import warnings
warnings.filterwarnings("ignore")

os.environ["TF_CPP_MIN_LOG_LEVEL"] = "2"
import tensorflow as tf
from sklearn.metrics import (
    accuracy_score, precision_score, recall_score, f1_score,
    confusion_matrix, classification_report
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# ── SETTINGS ──────────────────────────────────────────────────────────────────
MODEL_FOLDER  = "models"
RESULT_FOLDER = os.path.join("result", "confusion_matrix")
EXCEL_FILE    = os.path.join("result", "ALL_RESULTS.xlsx")
os.makedirs(RESULT_FOLDER, exist_ok=True)
os.makedirs(MODEL_FOLDER, exist_ok=True)
os.makedirs("result", exist_ok=True)

RUN_TIME = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

# ── STYLES ────────────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", start_color="2E4057", end_color="2E4057")
HEADER_FONT   = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL      = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
GREEN_FILL    = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
ORANGE_FILL   = PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")
RED_FILL      = PatternFill("solid", start_color="FFC7CE",  end_color="FFC7CE")
CENTER        = Alignment(horizontal="center", vertical="center")
THIN          = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)

def style_header_row(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = CENTER
        cell.border    = THIN

def style_data_row(ws, row, n_cols, alt=False):
    fill = ALT_FILL if alt else PatternFill()
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill      = fill
        cell.alignment = CENTER
        cell.border    = THIN

def color_metric_cell(cell, value):
    """Green if >=90%, Orange if >=70%, Red if <70%"""
    try:
        v = float(str(value).replace("%",""))
        if v >= 90:
            cell.fill = GREEN_FILL
        elif v >= 70:
            cell.fill = ORANGE_FILL
        else:
            cell.fill = RED_FILL
    except:
        pass

# ── LOAD MODELS ───────────────────────────────────────────────────────────────
def load_models_and_data():
    print("📂 Loading models and test data...")

    X_test = np.load(os.path.join(MODEL_FOLDER, "X_test.npy"))
    y_test = np.load(os.path.join(MODEL_FOLDER, "y_test.npy"))

    with open(os.path.join("processed", "label_encoder.pkl"), "rb") as f:
        le = pickle.load(f)

    models = {}
    for name in ["cnn_bilstm", "transformer"]:
        path = os.path.join(MODEL_FOLDER, f"{name}_best.keras")
        if os.path.exists(path):
            models[name] = tf.keras.models.load_model(path)
            print(f"   ✅ Loaded: {name}")

    if not models:
        raise FileNotFoundError("No trained models found! Run 3_train_classifier.py first.")

    return models, X_test, y_test, le

# ── EVALUATE ONE MODEL ────────────────────────────────────────────────────────
def evaluate_model(model, X_test, y_test, model_name, class_names):
    print(f"\n📊 Evaluating: {model_name}")

    y_pred_proba = model.predict(X_test, verbose=0)
    y_pred = np.argmax(y_pred_proba, axis=1)
    y_true = np.argmax(y_test,      axis=1)

    acc  = accuracy_score(y_true, y_pred)
    prec = precision_score(y_true, y_pred, average="weighted", zero_division=0)
    rec  = recall_score(y_true, y_pred, average="weighted", zero_division=0)
    f1   = f1_score(y_true, y_pred, average="weighted", zero_division=0)

    # Per-class metrics
    per_class_f1  = f1_score(y_true, y_pred, average=None, zero_division=0)
    per_class_prec = precision_score(y_true, y_pred, average=None, zero_division=0)
    per_class_rec  = recall_score(y_true, y_pred, average=None, zero_division=0)

    print(f"   Accuracy  : {acc*100:.2f}%")
    print(f"   Precision : {prec*100:.2f}%")
    print(f"   Recall    : {rec*100:.2f}%")
    print(f"   F1 Score  : {f1*100:.2f}%")

    metrics = {
        "run_time":  RUN_TIME,
        "model":     model_name,
        "accuracy":  round(acc  * 100, 2),
        "precision": round(prec * 100, 2),
        "recall":    round(rec  * 100, 2),
        "f1_score":  round(f1   * 100, 2),
    }

    per_class = []
    for i, name in enumerate(class_names):
        per_class.append({
            "run_time":   RUN_TIME,
            "model":      model_name,
            "activity":   name,
            "precision":  round(per_class_prec[i] * 100, 2) if i < len(per_class_prec) else 0,
            "recall":     round(per_class_rec[i]  * 100, 2) if i < len(per_class_rec)  else 0,
            "f1_score":   round(per_class_f1[i]   * 100, 2) if i < len(per_class_f1)   else 0,
        })

    return metrics, y_true, y_pred, per_class

# ── CONFUSION MATRIX PLOT ─────────────────────────────────────────────────────
def plot_confusion_matrix(y_true, y_pred, class_names, model_name):
    cm      = confusion_matrix(y_true, y_pred)
    cm_norm = cm.astype(float) / cm.sum(axis=1, keepdims=True)

    fig, axes = plt.subplots(1, 2, figsize=(20, 8))

    sns.heatmap(cm, annot=True, fmt="d", cmap="Blues",
                xticklabels=class_names, yticklabels=class_names, ax=axes[0])
    axes[0].set_title(f"{model_name} — Confusion Matrix (Count)",
                      fontsize=12, fontweight="bold")
    axes[0].set_xlabel("Predicted")
    axes[0].set_ylabel("True")
    axes[0].tick_params(axis="x", rotation=45)

    sns.heatmap(cm_norm, annot=True, fmt=".2f", cmap="Greens",
                xticklabels=class_names, yticklabels=class_names,
                ax=axes[1], vmin=0, vmax=1)
    axes[1].set_title(f"{model_name} — Confusion Matrix (Normalized)",
                      fontsize=12, fontweight="bold")
    axes[1].set_xlabel("Predicted")
    axes[1].set_ylabel("True")
    axes[1].tick_params(axis="x", rotation=45)

    plt.tight_layout()
    path = os.path.join(RESULT_FOLDER, f"{model_name}_confusion_matrix.png")
    plt.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    print(f"   📈 Confusion matrix saved → {path}")

# ── BUILD / UPDATE EXCEL ──────────────────────────────────────────────────────
def save_to_excel(all_summary, all_per_class, class_names):
    """
    Creates or updates ALL_RESULTS.xlsx with 5 sheets:

    Sheet 1 — Summary Table       (one row per model per run)
    Sheet 2 — Per-Class Metrics   (per activity breakdown)
    Sheet 3 — Model Comparison    (side-by-side best metrics)
    Sheet 4 — Confusion Matrix Info
    Sheet 5 — Hyperparameters
    """
    print(f"\n💾 Saving results to Excel: {EXCEL_FILE}")

    # ── Load existing or create new ──────────────────────────────────────────
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        print("   📂 Existing file found — appending new run...")
    else:
        wb = Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        print("   📄 Creating new Excel file...")

    # ═════════════════════════════════════════════════════════════════════
    # SHEET 1 — SUMMARY TABLE
    # ═════════════════════════════════════════════════════════════════════
    SHEET1 = "Summary_Table"
    if SHEET1 not in wb.sheetnames:
        ws1 = wb.create_sheet(SHEET1)
        headers = ["Run Time", "Model Name", "Accuracy (%)",
                   "Precision (%)", "Recall (%)", "F1 Score (%)"]
        ws1.append(headers)
        style_header_row(ws1, 1, len(headers))
        for col in [1,2]:
            ws1.column_dimensions[get_column_letter(col)].width = 22
        for col in [3,4,5,6]:
            ws1.column_dimensions[get_column_letter(col)].width = 16
    else:
        ws1 = wb[SHEET1]

    # Append new rows
    for i, m in enumerate(all_summary):
        row_idx = ws1.max_row + 1
        ws1.append([
            m["run_time"], m["model"],
            m["accuracy"], m["precision"], m["recall"], m["f1_score"]
        ])
        alt = (row_idx % 2 == 0)
        style_data_row(ws1, row_idx, 6, alt)
        # Color-code metric cells
        for col in [3, 4, 5, 6]:
            color_metric_cell(ws1.cell(row_idx, col),
                              ws1.cell(row_idx, col).value)

    # ═════════════════════════════════════════════════════════════════════
    # SHEET 2 — PER-CLASS METRICS
    # ═════════════════════════════════════════════════════════════════════
    SHEET2 = "Per_Class_Metrics"
    if SHEET2 not in wb.sheetnames:
        ws2 = wb.create_sheet(SHEET2)
        headers2 = ["Run Time", "Model", "Activity Class",
                    "Precision (%)", "Recall (%)", "F1 Score (%)"]
        ws2.append(headers2)
        style_header_row(ws2, 1, len(headers2))
        for col in [1,2,3]:
            ws2.column_dimensions[get_column_letter(col)].width = 24
        for col in [4,5,6]:
            ws2.column_dimensions[get_column_letter(col)].width = 16
    else:
        ws2 = wb[SHEET2]

    for pc in all_per_class:
        row_idx = ws2.max_row + 1
        ws2.append([
            pc["run_time"], pc["model"], pc["activity"],
            pc["precision"], pc["recall"], pc["f1_score"]
        ])
        alt = (row_idx % 2 == 0)
        style_data_row(ws2, row_idx, 6, alt)
        for col in [4, 5, 6]:
            color_metric_cell(ws2.cell(row_idx, col),
                              ws2.cell(row_idx, col).value)

    # ═════════════════════════════════════════════════════════════════════
    # SHEET 3 — MODEL COMPARISON (best of this run)
    # ═════════════════════════════════════════════════════════════════════
    SHEET3 = "Model_Comparison"
    if SHEET3 not in wb.sheetnames:
        ws3 = wb.create_sheet(SHEET3)
        h3 = ["Run Time", "Model", "Accuracy (%)", "Precision (%)",
              "Recall (%)", "F1 Score (%)", "Winner?"]
        ws3.append(h3)
        style_header_row(ws3, 1, len(h3))
        for col in range(1, 8):
            ws3.column_dimensions[get_column_letter(col)].width = 18
    else:
        ws3 = wb[SHEET3]

    best_f1 = max(all_summary, key=lambda x: x["f1_score"])
    for m in all_summary:
        row_idx = ws3.max_row + 1
        winner  = "🏆 BEST" if m["model"] == best_f1["model"] else ""
        ws3.append([
            m["run_time"], m["model"],
            m["accuracy"], m["precision"], m["recall"], m["f1_score"],
            winner
        ])
        alt = (row_idx % 2 == 0)
        style_data_row(ws3, row_idx, 7, alt)
        for col in [3, 4, 5, 6]:
            color_metric_cell(ws3.cell(row_idx, col),
                              ws3.cell(row_idx, col).value)
        if winner:
            ws3.cell(row_idx, 7).font = Font(bold=True, color="276221")

    # ═════════════════════════════════════════════════════════════════════
    # SHEET 4 — CONFUSION MATRIX INFO
    # ═════════════════════════════════════════════════════════════════════
    SHEET4 = "Confusion_Matrix_Info"
    if SHEET4 not in wb.sheetnames:
        ws4 = wb.create_sheet(SHEET4)
        h4 = ["Run Time", "Model", "Image Saved At",
              "True Positive Rate (avg)", "Notes"]
        ws4.append(h4)
        style_header_row(ws4, 1, len(h4))
        for col in range(1, 6):
            ws4.column_dimensions[get_column_letter(col)].width = 32
    else:
        ws4 = wb[SHEET4]

    for m in all_summary:
        row_idx = ws4.max_row + 1
        img_path = os.path.join(RESULT_FOLDER, f"{m['model']}_confusion_matrix.png")
        ws4.append([
            m["run_time"], m["model"], img_path,
            f"{m['recall']}%",
            "See confusion matrix image in result/confusion_matrix/"
        ])
        style_data_row(ws4, row_idx, 5, (row_idx % 2 == 0))

    # ═════════════════════════════════════════════════════════════════════
    # SHEET 5 — HYPERPARAMETERS
    # ═════════════════════════════════════════════════════════════════════
    SHEET5 = "Hyperparameters"
    if SHEET5 not in wb.sheetnames:
        ws5 = wb.create_sheet(SHEET5)
        h5 = ["Parameter", "Value", "Description"]
        ws5.append(h5)
        style_header_row(ws5, 1, 3)
        ws5.column_dimensions["A"].width = 28
        ws5.column_dimensions["B"].width = 20
        ws5.column_dimensions["C"].width = 40

        params = [
            ("Window Size",        "128",          "Number of sensor readings per sample"),
            ("Step Size",          "64",           "Overlap between windows (50%)"),
            ("Batch Size",         "64",           "Samples per training step"),
            ("Epochs (max)",       "50",           "Maximum training rounds"),
            ("Learning Rate",      "0.001",        "Initial Adam optimizer LR"),
            ("Optimizer",          "Adam",         "Adaptive Moment Estimation"),
            ("Loss Function",      "Categorical Crossentropy", "For multi-class classification"),
            ("Activation (hidden)","ReLU",         "Rectified Linear Unit"),
            ("Activation (output)","Softmax",      "Outputs probability per class"),
            ("Train Split",        "80%",          "80% training, 20% testing"),
            ("Val Split",          "10%",          "10% of train used for validation"),
            ("EarlyStopping",      "Patience=8",   "Stop if val_loss not improving"),
            ("ReduceLROnPlateau",  "Factor=0.5, Patience=4", "Halve LR if stuck"),
            ("CNN Filters",        "64→128→256",   "Increasing feature maps"),
            ("BiLSTM Units",       "128→64",       "Bidirectional LSTM hidden units"),
            ("Transformer Heads",  "4",            "Multi-head attention heads"),
            ("Key Dimension",      "32",           "Attention key dimension"),
            ("Dropout Rate",       "0.3–0.4",      "Regularization to prevent overfitting"),
            ("Class Weights",      "Balanced",     "Handle class imbalance"),
            ("Number of Classes",  "12",           "Activity types in HuGaDB"),
        ]
        for i, (param, val, desc) in enumerate(params):
            ws5.append([param, val, desc])
            row_idx = i + 2
            style_data_row(ws5, row_idx, 3, (row_idx % 2 == 0))
            ws5.cell(row_idx, 1).font = Font(bold=True)

    wb.save(EXCEL_FILE)
    print(f"   ✅ Excel saved → {EXCEL_FILE}")
    print(f"\n   📋 Sheets in Excel:")
    for s in wb.sheetnames:
        print(f"      - {s}")

# ── MAIN ──────────────────────────────────────────────────────────────────────
def main():
    print("=" * 60)
    print("  GAIT PROJECT — Step 5: Evaluation")
    print(f"  Run Time: {RUN_TIME}")
    print("=" * 60)

    models, X_test, y_test, le = load_models_and_data()
    class_names = list(le.classes_)

    all_summary   = []
    all_per_class = []

    for model_name, model in models.items():
        metrics, y_true, y_pred, per_class = evaluate_model(
            model, X_test, y_test, model_name, class_names
        )
        all_summary.append(metrics)
        all_per_class.extend(per_class)
        plot_confusion_matrix(y_true, y_pred, class_names, model_name)

    # Save everything to Excel
    save_to_excel(all_summary, all_per_class, class_names)

    # Final summary
    best = max(all_summary, key=lambda x: x["f1_score"])
    print("\n" + "=" * 60)
    print("  ✅ EVALUATION COMPLETE!")
    print(f"  🏆 Best Model : {best['model']}")
    print(f"     Accuracy   : {best['accuracy']}%")
    print(f"     Precision  : {best['precision']}%")
    print(f"     Recall     : {best['recall']}%")
    print(f"     F1 Score   : {best['f1_score']}%")
    print(f"\n  📊 Excel file : {EXCEL_FILE}")
    print(f"     → Sheet 1  : Summary Table (overall metrics)")
    print(f"     → Sheet 2  : Per-Class Metrics (per activity)")
    print(f"     → Sheet 3  : Model Comparison (winner highlighted)")
    print(f"     → Sheet 4  : Confusion Matrix Info")
    print(f"     → Sheet 5  : Hyperparameters Table")
    print("=" * 60)
    print("\n▶ Next step: python predict.py")

if __name__ == "__main__":
    main()
